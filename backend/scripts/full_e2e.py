from __future__ import annotations

import argparse
import json
import os
import tempfile
import time
from pathlib import Path

import httpx
from docx import Document
from PIL import Image, ImageDraw

from smoke_e2e import expect, sign_in, wait_for_job


def create_project(client: httpx.Client) -> dict:
    stamp = int(time.time())
    return expect(
        client.post(
            "/projects",
            json={
                "name": f"Full E2E {stamp}",
                "description": "Disposable full-feature verification project",
                "target_score": 90,
                "exam_date": None,
                "weekly_study_minutes": 300,
            },
        ),
        201,
    )


def make_docx(path: Path, revision: int) -> None:
    document = Document()
    document.add_heading("Transformer Attention", level=1)
    document.add_paragraph(
        "Attention trong Transformer sử dụng Query, Key và Value. "
        "Scaled dot-product attention so sánh Query với Key, dùng softmax "
        "để tạo trọng số rồi tổng hợp các vector Value. "
        f"Đây là bản tài liệu kiểm thử số {revision}."
    )
    document.add_heading("Ứng dụng", level=2)
    document.add_paragraph(
        "Self-attention giúp mô hình học quan hệ xa trong chuỗi và xử lý song song."
    )
    document.save(path)


def upload_document(
    client: httpx.Client,
    project_id: str,
    path: Path,
    mime_type: str,
) -> dict:
    with path.open("rb") as handle:
        return expect(
            client.post(
                f"/projects/{project_id}/documents",
                files={"file": (path.name, handle, mime_type)},
            ),
            202,
        )


def upload_version(
    client: httpx.Client,
    project_id: str,
    document_id: str,
    path: Path,
) -> dict:
    with path.open("rb") as handle:
        return expect(
            client.post(
                f"/projects/{project_id}/documents/{document_id}/versions",
                files={
                    "file": (
                        path.name,
                        handle,
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                },
            ),
            202,
        )


def document_detail(
    client: httpx.Client,
    project_id: str,
    document_id: str,
) -> dict:
    return expect(
        client.get(f"/projects/{project_id}/documents/{document_id}"),
        200,
    )


def wait_for_active_version(
    client: httpx.Client,
    project_id: str,
    document_id: str,
    expected_version_id: str,
) -> dict:
    detail = document_detail(client, project_id, document_id)
    if detail.get("active_version_id") != expected_version_id:
        raise RuntimeError(
            "Unexpected active version after ingestion: "
            + json.dumps(detail, ensure_ascii=False)
        )
    version = next(
        (row for row in detail.get("versions", []) if row["id"] == expected_version_id),
        None,
    )
    if not version or version.get("status") != "ready":
        raise RuntimeError("Expected document version to be ready")
    return detail


def make_answer_image(path: Path) -> None:
    image = Image.new("RGB", (1100, 360), "white")
    draw = ImageDraw.Draw(image)
    draw.text(
        (40, 60),
        "Query, Key and Value are the three vectors used by attention.",
        fill="black",
    )
    draw.text(
        (40, 150),
        "Softmax creates weights used to combine the Value vectors.",
        fill="black",
    )
    image.save(path, format="PNG")


def check_profile_and_project(
    client: httpx.Client,
    project: dict,
) -> dict:
    me = expect(client.get("/auth/me"), 200)
    updated_me = expect(
        client.patch(
            "/auth/me",
            json={"full_name": "Codex Full E2E", "timezone": "Asia/Ho_Chi_Minh"},
        ),
        200,
    )
    if updated_me.get("timezone") != "Asia/Ho_Chi_Minh":
        raise RuntimeError("Profile timezone update did not persist")

    updated_project = expect(
        client.patch(
            f"/projects/{project['id']}",
            json={
                "name": project["name"] + " Updated",
                "description": "Updated during disposable full E2E",
                "target_score": 92,
                "weekly_study_minutes": 360,
            },
        ),
        200,
    )
    if updated_project.get("target_score") != 92:
        raise RuntimeError("Project update did not persist")
    print(f"[profile-project] user={me['user_id']} project updated")
    return updated_project


def check_chat(
    client: httpx.Client,
    project_id: str,
) -> None:
    session = expect(
        client.post(f"/projects/{project_id}/chat/sessions"),
        201,
    )
    session_id = session["id"]
    answer = expect(
        client.post(
            f"/projects/{project_id}/chat/sessions/{session_id}/messages",
            json={"content": "Attention sử dụng ba vector nào?"},
        ),
        201,
    )
    if not answer.get("citations"):
        raise RuntimeError("Non-streaming RAG answer has no citation")

    stream_session = expect(
        client.post(f"/projects/{project_id}/chat/sessions"),
        201,
    )
    response = client.post(
        f"/projects/{project_id}/chat/sessions/{stream_session['id']}/messages/stream",
        json={"content": "Query, Key và Value có vai trò gì?"},
        timeout=90,
    )
    if response.status_code != 200:
        raise RuntimeError(f"Streaming chat failed: {response.status_code} {response.text[:500]}")
    if "event: error" in response.text:
        raise RuntimeError("Streaming chat emitted an error event: " + response.text[-1000:])
    if "event: done" not in response.text:
        raise RuntimeError("Streaming chat did not emit a done event")
    messages = expect(
        client.get(
            f"/projects/{project_id}/chat/sessions/{stream_session['id']}/messages"
        ),
        200,
    )
    if len(messages) < 2:
        raise RuntimeError("Streaming chat messages were not persisted")
    print("[chat] regular and streaming RAG passed")


def check_pdf_annotations(
    client: httpx.Client,
    project_id: str,
    document_id: str,
    version_id: str,
) -> None:
    signed = expect(client.get(f"/document-versions/{version_id}/signed-url"), 200)
    if not str(signed.get("signed_url") or "").startswith("http"):
        raise RuntimeError("Document signed URL is missing")

    annotation = expect(
        client.post(
            f"/projects/{project_id}/documents/{document_id}/versions/{version_id}/annotations",
            json={
                "page_number": 1,
                "annotation_type": "rectangle",
                "rectangles": [{"x": 0.1, "y": 0.1, "width": 0.3, "height": 0.1}],
                "content": "Disposable E2E note",
                "color": "#f4d06f",
            },
        ),
        201,
    )
    rows = expect(
        client.get(
            f"/projects/{project_id}/documents/{document_id}/versions/{version_id}/annotations",
            params={"page": 1},
        ),
        200,
    )
    if not any(row["id"] == annotation["id"] for row in rows):
        raise RuntimeError("Created annotation was not listed")
    updated = expect(
        client.patch(
            f"/annotations/{annotation['id']}",
            json={"content": "Updated disposable note", "color": "#00aa55"},
        ),
        200,
    )
    if updated.get("color") != "#00AA55" or updated.get("version") != 2:
        raise RuntimeError("Annotation update did not persist")
    expect(client.delete(f"/annotations/{annotation['id']}"), 204)
    print("[annotations] signed URL and CRUD passed")


def check_roadmap_and_schedules(
    client: httpx.Client,
    project_id: str,
) -> None:
    generated = expect(
        client.post(f"/projects/{project_id}/roadmap/generate", timeout=120),
        201,
    )
    if not generated.get("topics"):
        raise RuntimeError("Roadmap generation returned no topics")
    roadmap = expect(client.get(f"/projects/{project_id}/roadmap"), 200)
    if not roadmap.get("topics"):
        raise RuntimeError("Roadmap retrieval returned no topics")

    schedules = expect(
        client.post(f"/projects/{project_id}/schedules/generate", timeout=60),
        201,
    )
    if len(schedules) < 2:
        raise RuntimeError("Schedule generation returned too few entries")
    accepted = expect(client.post(f"/schedules/{schedules[0]['id']}/accept"), 200)
    if accepted.get("suggestion_status") != "accepted":
        raise RuntimeError("Schedule accept did not persist")
    expect(client.post(f"/schedules/{schedules[0]['id']}/complete"), 201)
    listed = expect(client.get(f"/projects/{project_id}/schedules"), 200)
    selected = next(row for row in listed if row["id"] == schedules[0]["id"])
    if selected.get("completed_by_me") is not True:
        raise RuntimeError("Schedule completion was not listed")
    expect(client.delete(f"/schedules/{schedules[0]['id']}/complete"), 204)
    rejected = expect(client.post(f"/schedules/{schedules[1]['id']}/reject"), 200)
    if rejected.get("suggestion_status") != "rejected":
        raise RuntimeError("Schedule reject did not persist")
    print("[roadmap-schedule] generation and state transitions passed")


def check_quiz_and_ocr(
    client: httpx.Client,
    project_id: str,
    answer_image: Path,
) -> None:
    mcq = expect(
        client.post(
            f"/projects/{project_id}/quiz/generate",
            json={"count": 1, "question_type": "mcq"},
            timeout=120,
        ),
        201,
    )
    mcq_question = mcq["questions"][0]
    attempt = expect(
        client.post(
            f"/projects/{project_id}/quiz/sessions/{mcq['session']['id']}/questions/{mcq_question['id']}/answer",
            json={"answer": mcq_question["correct_answer"]},
            timeout=90,
        ),
        200,
    )
    if attempt.get("is_correct") is not True:
        raise RuntimeError("Correct MCQ answer was not graded correctly")
    expect(
        client.post(
            f"/projects/{project_id}/quiz/sessions/{mcq['session']['id']}/submit"
        ),
        200,
    )
    expect(client.get(f"/projects/{project_id}/quiz/sessions"), 200)
    expect(
        client.get(
            f"/projects/{project_id}/quiz/sessions/{mcq['session']['id']}"
        ),
        200,
    )
    bank = expect(
        client.post(
            f"/projects/{project_id}/quiz/start",
            json={"count": 1, "question_type": "mcq"},
        ),
        201,
    )
    if not bank.get("questions"):
        raise RuntimeError("Starting quiz from bank returned no questions")

    essay = expect(
        client.post(
            f"/projects/{project_id}/quiz/generate",
            json={"count": 1, "question_type": "essay"},
            timeout=120,
        ),
        201,
    )
    essay_question = essay["questions"][0]
    with answer_image.open("rb") as handle:
        scanned = expect(
            client.post(
                f"/projects/{project_id}/quiz/sessions/{essay['session']['id']}/questions/{essay_question['id']}/scan",
                files={"file": (answer_image.name, handle, "image/png")},
                timeout=120,
            ),
            201,
        )
    if scanned.get("status") != "ocr_pending_confirmation":
        raise RuntimeError("OCR scan did not enter confirmation state")
    signed = expect(client.get(f"/quiz-attempts/{scanned['id']}/scan-url"), 200)
    if not signed.get("signed_url"):
        raise RuntimeError("OCR scan signed URL is missing")
    confirmed_text = scanned.get("ocr_raw_text") or (
        "Query, Key and Value are combined with softmax attention weights."
    )
    confirmed = expect(
        client.post(
            f"/quiz-attempts/{scanned['id']}/confirm-scan",
            json={"text": confirmed_text},
            timeout=120,
        ),
        200,
    )
    if confirmed.get("status") != "graded":
        raise RuntimeError("Confirmed OCR answer was not graded")
    expect(client.delete(f"/quiz-attempts/{scanned['id']}/scan"), 204)
    expect(
        client.post(
            f"/projects/{project_id}/quiz/sessions/{essay['session']['id']}/submit"
        ),
        200,
    )
    expect(client.get(f"/projects/{project_id}/reviews/due"), 200)
    print("[quiz-ocr] MCQ, bank, OCR confirmation and essay grading passed")


def check_progress_and_search(
    client: httpx.Client,
    project_id: str,
    project_name: str,
) -> None:
    results = expect(client.get("/search", params={"q": "Full E2E"}), 200)
    if not any(row.get("project_id") == project_id for row in results):
        raise RuntimeError("Global search did not return the project")
    overview = expect(client.get("/progress/overview"), 200)
    if overview.get("projects", 0) < 1 or overview.get("quiz_attempts", 0) < 1:
        raise RuntimeError("Progress overview did not aggregate E2E activity")
    rebuilt = expect(
        client.post(
            f"/projects/{project_id}/progress/rebuild",
            params={"days": 7},
            timeout=120,
        ),
        200,
    )
    if rebuilt.get("snapshots_rebuilt", 0) < 1:
        raise RuntimeError("Progress rebuild produced no snapshots")
    history = expect(
        client.get("/progress/history", params={"days": 7, "project_id": project_id}),
        200,
    )
    if not history:
        raise RuntimeError("Progress history is empty after rebuild")
    print(f"[progress-search] project={project_name} aggregation passed")


def cleanup_project(client: httpx.Client, project_id: str) -> None:
    response = client.delete(f"/projects/{project_id}")
    if response.status_code not in {204, 404}:
        raise RuntimeError(
            f"Project cleanup failed: {response.status_code} {response.text[:500]}"
        )


def main() -> None:
    parser = argparse.ArgumentParser(description="Run comprehensive RAGTutor E2E checks.")
    parser.add_argument(
        "--base-url",
        default=os.getenv("RAGTUTOR_API_URL", "http://127.0.0.1:8000/api/v1"),
    )
    parser.add_argument("--email", default=os.getenv("RAGTUTOR_TEST_EMAIL", ""))
    parser.add_argument("--password", default=os.getenv("RAGTUTOR_TEST_PASSWORD", ""))
    parser.add_argument("--pdf", type=Path, default=Path("learning/sample_rag.pdf"))
    parser.add_argument("--timeout", type=int, default=300)
    args = parser.parse_args()

    if not args.email or not args.password:
        raise SystemExit("Provide disposable E2E credentials")
    if not args.pdf.is_file():
        raise SystemExit(f"PDF fixture does not exist: {args.pdf}")

    project_id: str | None = None
    with httpx.Client(base_url=args.base_url.rstrip("/"), timeout=60) as client:
        token = sign_in(client, args.email, args.password)
        client.headers["Authorization"] = f"Bearer {token}"
        project = create_project(client)
        project_id = project["id"]
        print(f"[project] created {project_id}")

        try:
            updated_project = check_profile_and_project(client, project)
            with tempfile.TemporaryDirectory(prefix="ragtutor-full-e2e-") as temp_dir:
                temp = Path(temp_dir)
                doc_v1 = temp / "attention_v1.docx"
                doc_v2 = temp / "attention_v2.docx"
                scan = temp / "essay_answer.png"
                make_docx(doc_v1, 1)
                make_docx(doc_v2, 2)
                make_answer_image(scan)

                doc_upload = upload_document(
                    client,
                    project_id,
                    doc_v1,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
                wait_for_job(client, doc_upload["job_id"], args.timeout)
                wait_for_active_version(
                    client,
                    project_id,
                    doc_upload["document_id"],
                    doc_upload["version_id"],
                )

                version_upload = upload_version(
                    client,
                    project_id,
                    doc_upload["document_id"],
                    doc_v2,
                )
                wait_for_job(client, version_upload["job_id"], args.timeout)
                wait_for_active_version(
                    client,
                    project_id,
                    doc_upload["document_id"],
                    version_upload["version_id"],
                )
                print("[documents] upload and immutable version activation passed")

                pdf_upload = upload_document(
                    client,
                    project_id,
                    args.pdf,
                    "application/pdf",
                )
                wait_for_job(client, pdf_upload["job_id"], args.timeout)
                wait_for_active_version(
                    client,
                    project_id,
                    pdf_upload["document_id"],
                    pdf_upload["version_id"],
                )
                expect(client.get(f"/projects/{project_id}/documents"), 200)
                expect(client.get(f"/projects/{project_id}/document-jobs"), 200)

                check_pdf_annotations(
                    client,
                    project_id,
                    pdf_upload["document_id"],
                    pdf_upload["version_id"],
                )
                check_chat(client, project_id)
                check_roadmap_and_schedules(client, project_id)
                check_quiz_and_ocr(client, project_id, scan)
                check_progress_and_search(
                    client,
                    project_id,
                    updated_project["name"],
                )

                deletion = expect(
                    client.delete(
                        f"/projects/{project_id}/documents/{pdf_upload['document_id']}"
                    ),
                    202,
                )
                wait_for_job(client, deletion["job_id"], args.timeout)
                print("[documents] asynchronous permanent deletion passed")

            expect(client.post("/auth/signout"), 204)
            print(json.dumps({"status": "PASS", "project_id": project_id}, indent=2))
        finally:
            cleanup_project(client, project_id)
            print(f"[cleanup] deleted temporary project {project_id}")


if __name__ == "__main__":
    main()
